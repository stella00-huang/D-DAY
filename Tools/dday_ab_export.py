"""Convert the authoritative AB workbook into Unity JsonUtility data without editing Excel."""

import argparse
import hashlib
import json
import re
import random
from pathlib import Path

import openpyxl

from dday_ab_solver import NAMES, SEATS, TUNED, parse_day, solve, satisfied


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:/Users/stella/Desktop/配置表AB.xlsx")
TARGET = ROOT / "Assets/Resources/DDay/ABConfig.json"
IDS = {"SISY": "rabbit", "ORIN": "wolf", "NARCIS": "fox",
       "ECHO": "shiba", "METIS": "cat", "THANE": "bear", "CASSIA": "panda"}
RUNTIME_ORDER = ("SISY", "ORIN", "NARCIS", "ECHO", "METIS", "THANE", "CASSIA")
TASK_IDS = {4: "report", 5: "legacy", 6: "orin_inquiry", 7: "hr",
            8: "coordination", 9: "archive", 10: "promotion",
            11: "belongings", 12: "revised_report", 13: "recruitment",
            14: "day6_wrap", 15: "report", 16: "unfinished",
            17: "backlog", 18: "hr", 19: "coordination",
            20: "archive", 21: "promotion", 22: "history",
            23: "revised_report", 24: "recruitment", 25: "day6_wrap", 26: "no_b1", 27: "no_b1"}


def value(row, col):
    return (row[col - 1].value or "").strip()


def loop_day(row):
    loop = 1 if value(row, 1) == "一周目" else 2 if value(row, 1) == "二周目" else 0
    day = int(re.search(r"Day\s*(\d+)", value(row, 2)).group(1))
    return loop, day


def rules_from_day(parsed):
    result = []
    for item in parsed["rules"]:
        kind, target = item["kind"], item["target"]
        if kind == "seat":
            kind, target = "seat_is", str(SEATS.index(target))
        elif kind == "zone":
            kind, target = "seat_tag", {"boss": "near_boss",
                "printer": "near_printer", "lounge": "near_lounge",
                "coffee": "near_coffee", "far_coffee": "far_coffee",
                "far_boss": "far_boss"}.get(target, target)
        elif kind == "near_seat":
            target = "rabbit_seat"
        elif kind == "not_adjacent_seat":
            kind, target = "not_adjacent", "rabbit_seat"
        else:
            target = IDS.get(target, target)
        result.append({"animal": IDS[item["actor"]], "kind": kind,
                       "target": target, "distance": 3 if kind == "far_person"
                       else 2 if kind in ("near_person", "near_seat") else 0,
                       "text": item["source"], "group": "A1"})
    return result


def initial_layout(day):
    """A reproducible shuffled preset, independently checked against every A1 clause.

    The solver's example answer is never used as the starting arrangement.
    """
    rng = random.Random(20261011 + day["loop"] * 100 + day["day"])
    movable = [name for name in day["roster"]
               if not (day["loop"] == 1 and day["day"] == 1 and name == "SISY")]
    available = [i for i in range(8) if not (day["loop"] == 1 and i == 6)]
    for attempt in range(1000):
        rng.shuffle(available)
        positions = dict(zip(movable, available))
        if day["loop"] == 1 and day["day"] == 1:
            positions["SISY"] = 6
        unmet = sum(not satisfied(rule, positions, TUNED) for rule in day["rules"])
        if unmet >= 3:
            return [positions.get(name, -1) for name in RUNTIME_ORDER], unmet
    raise ValueError(f"Cannot generate unsolved initial layout: {day['loop']}/{day['day']}")


def dialogue_lines(raw):
    # Preserve every word; only split explicit inline turns and blank separators.
    return [part.strip() for line in raw.splitlines() if line.strip()
            for part in re.split(r"(?<=。)\s+(?=(?:" + "|".join(NAMES) + r"|老板)[：:])", line)
            if part.strip()]


def task_branches(raw, row_index):
    branches = []
    key, body, other = [], [], False
    def flush():
        if key:
            text = "\n".join(body).strip()
            for name in key:
                branches.append({"owner": IDS[name], "text": text, "narration": False,
                    "condition": "sisy_unassigned" if row_index == 15 and name in ("SISY", "ORIN") else ""})
    for line in raw.splitlines():
        stripped = line.strip()
        header = re.match(r"【关键｜([^】]+)】(.*)", stripped)
        if header:
            flush(); key = [name for name in NAMES if name in header.group(1)]
            body = []; other = False
        elif "其他角色（带过）" in stripped or "本任务无关键线" in stripped:
            flush(); key = []; body = []; other = True
        elif stripped.startswith("━━"):
            continue
        elif stripped.startswith("【条件】"):
            # Explicit authored alternative, never append the failed-state design note to dialogue.
            quote = re.search(r'“(.*?)”', stripped)
            if row_index == 15 and quote:
                branches.append({"owner": "rabbit", "text": "SISY：“"+quote.group(1)+"”",
                                 "narration": False, "condition": "sisy_assigned"})
        elif stripped.startswith("——不让任何人说出"):
            # Stage direction embedded in the story cell; retained only in WorkbookFull/audit.
            continue
        elif other:
            match = re.match(r"^("+"|".join(NAMES)+r")[：:]", stripped)
            if match:
                branches.append({"owner": IDS[match.group(1)], "text": stripped,
                                 "narration": True, "condition": ""})
            elif stripped:
                if not branches: raise ValueError(f"B1 row {row_index}: unmapped {stripped}")
                branches[-1]["text"] += "\n" + line
        elif key:
            body.append(line)
        elif stripped:
            raise ValueError(f"B1 row {row_index}: unmapped {stripped}")
    flush()
    return branches


def parse_tasks(sheet):
    tasks = []
    for row in sheet.iter_rows(min_row=4):
        index = row[0].row
        if not value(row, 1): continue
        loop, day = loop_day(row)
        if index not in TASK_IDS: raise ValueError(f"Unmapped B1 task row {index}")
        title = value(row, 3)
        branches = [] if day == 7 else task_branches(value(row, 5), index)
        tasks.append({"id": TASK_IDS[index], "loop": loop, "day": day,
            "title": title, "description": "",  # No public description column exists in B1.
            "branches": branches, "generic": "已完成分配的工作。",
            "fixed": "【固定】" in title, "sourceRow": index})
    return tasks


# User confirmed 2026-10-12: fill these four empty record cells from their existing scene.
RECORD_SUPPLEMENTS = {
    12: "ORIN 的包上挂着 SISY 送的旧兔脚护身符。它一直带着，也不让别人碰。",
    20: "ORIN 主动留意 SISY 有没有吃午饭。SISY 嘴上嫌它管得多，最后还是答应照顾自己。",
    21: "SISY 终于把找了很久的 ORIN 大学旧照送了出去。总拍别人的灰狼，这次也有了自己的照片。",
    39: "CASSIA 对 ORIN 的工作速度赞不绝口，ORIN 却记得它昨天摸了一下午鱼。夸年轻动物之前，先藏好摸鱼记录吧。"
}


def player_clue_text(row, text):
    if row == 14:
        return text.split("\n不再用它承担")[0]
    if row == 36:
        return text.split("此时只确认它存在")[0]
    return text


def parse_clues(sheet):
    # Authored A2 locations are heterogeneous prose. Normalize their trigger
    # without changing any source text. Index is the authoritative worksheet row.
    conditions = {
        4: "adjacent", 5: "adjacent", 6: "center", 7: "seat_b3",
        8: "edge_not_b3", 9: "printer", 10: "adjacent", 11: "boss",
        12: "desk", 13: "boss", 14: "desk", 15: "adjacent",
        16: "desk", 17: "lounge", 18: "desk_or_lounge", 19: "near",
        20: "adjacent", 21: "adjacent", 22: "not_adjacent_no_sisy_task",
        23: "not_coffee", 24: "adjacent", 25: "printer", 26: "either_boss",
        27: "adjacent", 28: "adjacent_or_printer", 29: "after_promotion",
        30: "printer", 31: "desk", 32: "boss", 33: "desk",
        34: "bear_center", 35: "adjacent", 36: "desk_x_prereq",
        37: "adjacent", 38: "adjacent", 39: "adjacent_or_lounge",
        40: "lounge", 41: "desk", 42: "desk", 43: "center_or_desk",
    }
    clues = []
    for row in sheet.iter_rows(min_row=4):
        raw_loop = value(row, 1)
        loop = 0 if raw_loop == "两周目" else 1 if raw_loop == "一周目" else 2
        match = re.search(r"Day\s*(\d+)(?:[–-](\d+))?", value(row, 2))
        first = int(match.group(1))
        last = int(match.group(2) or first)
        location = value(row, 6)
        actor_text = value(row, 5)
        actors = [IDS[name] for name in NAMES if name in actor_text]
        if not actors:
            raise ValueError(f"A2 row {row[0].row}: no actor")
        if row[0].row == 29:  # Explicitly after the Day 4 evening promotion.
            first = last = 5
        clues.append({"id": f"A2_R{row[0].row:02d}", "loop": loop,
            "firstDay": first, "lastDay": last,
            "category": value(row, 3), "title": value(row, 4),
            "actors": actors, "location": location,
            "condition": conditions[row[0].row],
            "requiredEvent": {15: "B2_L1_D4", 26: "B2_L2_D2",
                29: "B2_L2_D4", 32: "B2_L2_D4", 33: "B2_L2_D4"}.get(row[0].row, ""),
            "stateText": value(row, 7), "description": value(row, 8) or RECORD_SUPPLEMENTS.get(row[0].row, ""),
            "dialogue": dialogue_lines(player_clue_text(row[0].row, value(row, 7))),
            "recordSupplement": not bool(value(row, 8)),
            "sourceRow": row[0].row})
    return clues


def split_event_text(raw):
    result = []
    for line in raw.splitlines():
        line = line.strip()
        if not line:
            continue
        # Preserve original wording and punctuation. Only paginate long prose.
        if len(line) <= 78:
            result.append(line)
            continue
        pieces = re.split(r"(?<=[。！？])", line)
        page = ""
        for piece in pieces:
            if page and len(page + piece) > 78:
                result.append(page)
                page = ""
            page += piece
        if page:
            result.append(page)
    return result


def comic_texts(loop, day, raw):
    if loop != 1 or day not in (1, 2, 4): return None
    paragraphs = [line for line in raw.splitlines() if line.strip()]
    def body(line): return re.sub(r"^[①-⑨][^：]*：", "", line).strip()
    if day == 1:
        collapse, photo = paragraphs[2].split("而口袋里的照片", 1)
        pages = [[body(paragraphs[0])], [body(paragraphs[1])], [collapse],
                 ["而口袋里的照片" + photo], [body(paragraphs[3])], [], []]
    elif day == 2:
        parts = re.findall(r"[^。]+。?", body(paragraphs[0]))
        pages = [[], [parts[0]], [], parts[1:], [body(paragraphs[1])], [], []]
    else:
        promotion, recruitment = body(paragraphs[0]).split("，与此同时", 1)
        pages = [[], [], [promotion + "，"], [], ["与此同时" + recruitment], [], [body(paragraphs[1])]]
    return [{"lines": lines} for lines in pages]


def parse_fixed(sheet):
    events = []
    for row in sheet.iter_rows(min_row=4):
        loop, day = loop_day(row)
        raw = value(row, 5)
        events.append({"id": f"B2_L{loop}_D{day}", "loop": loop,
            "day": day, "title": value(row, 4), "trigger": value(row, 3),
            "lines": split_event_text(raw), "original": raw,
            "comicTexts": comic_texts(loop, day, raw),
            "sourceRow": row[0].row})
    return events


def main():
    global SOURCE
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", type=Path, default=SOURCE)
    args = parser.parse_args()
    SOURCE = args.workbook
    workbook = openpyxl.load_workbook(SOURCE, data_only=False)
    assert workbook.sheetnames == ["A1-座位通关", "A2-即时线索",
                                  "B1-每日任务与剧情", "B2-固定事件"]
    seat_days = [parsed for row in workbook.worksheets[0].iter_rows(min_row=4)
                 if (parsed := parse_day(row)) is not None]
    if len(seat_days) != 12 or any(solve(day, TUNED)["count"] == 0 for day in seat_days):
        raise ValueError("A1 has an unsolved day; refusing formal export")
    days = [{"loop": day["loop"], "day": day["day"],
             "present": [IDS[name] for name in day["roster"]],
             "rules": rules_from_day(day), "sourceRow": day["row"],
             "initialSeats": initial_layout(day)[0],
             "initialUnmet": initial_layout(day)[1]}
            for day in seat_days]
    clues = parse_clues(workbook.worksheets[1])
    tasks = parse_tasks(workbook.worksheets[2])
    fixed = parse_fixed(workbook.worksheets[3])
    if len(clues) != 40 or len(tasks) != 24 or len(fixed) != 10:
        raise ValueError((len(clues), len(tasks), len(fixed)))
    config = {"sourceSha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              "seatDays": days, "clues": clues, "tasks": tasks,
              "fixedEvents": fixed}
    TARGET.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {TARGET}: {len(days)} seating days, {len(clues)} optional clues, "
          f"{len(tasks)} B1 rows, {len(fixed)} B2 events")


if __name__ == "__main__":
    main()

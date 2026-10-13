"""Read every A1 row in 配置表AB.xlsx and exhaustively solve the 2x4 desk puzzle.

This is an audit tool. It does not edit the workbook or Unity assets.
"""

import argparse
import hashlib
import itertools
import json
import re
from pathlib import Path

import openpyxl


NAMES = ("SISY", "ORIN", "NARCIS", "ECHO", "METIS", "CASSIA", "THANE")
SEATS = ("A1", "A2", "A3", "A4", "B1", "B2", "B3", "B4")
COORDS = tuple((i // 4, i % 4) for i in range(8))
BASELINE = {
    "edge": (0, 3, 4, 7), "center": (1, 2, 5, 6),
    "boss": (6, 7), "printer": (2, 3, 6),
    "lounge": (4, 5), "coffee": (0, 1, 2, 3),
    "far_boss": (0, 1, 2, 4), "far_coffee": (4, 7),
}
# Minimal reasonable area changes needed for the supplied A1 rules:
# B2 is within two desk steps of the lower-right boss office;
# B3 is within two desk steps of the left lounge. Neither facility moves.
TUNED = {**BASELINE, "boss": (5, 6, 7), "lounge": (4, 5, 6)}


def parse_clause(clause):
    clause = clause.strip().strip("。;； ")
    if not clause or clause in ("不可移动", "不固定B3"):
        return []
    if clause == "固定在B3":
        return [("seat", "B3")]
    if clause in ("位于边缘位置", "靠近边缘"):
        return [("zone", "edge")]
    if clause == "靠近中心":
        return [("zone", "center")]
    if clause in ("靠近老板", "靠近老板办公室"):
        return [("zone", "boss")]
    if clause == "远离老板":
        return [("zone", "far_boss")]
    if clause in ("靠近打印机", "靠近打印区"):
        return [("zone", "printer")]
    if clause == "靠近咖啡机":
        return [("zone", "coffee")]
    if clause == "远离咖啡机":
        return [("zone", "far_coffee")]
    if clause == "靠近休息区":
        return [("zone", "lounge")]
    if clause == "不与B3相邻":
        return [("not_adjacent_seat", "B3")]
    if clause == "靠近B3":
        return [("near_seat", "B3")]
    m = re.fullmatch(r"(?:与)?(" + "|".join(NAMES) + r")相邻", clause)
    if m:
        return [("adjacent", m.group(1))]
    m = re.fullmatch(r"靠近(" + "|".join(NAMES) + r")和(" + "|".join(NAMES) + r")", clause)
    if m:
        return [("near_person", m.group(1)), ("near_person", m.group(2))]
    m = re.fullmatch(r"靠近(" + "|".join(NAMES) + r")", clause)
    if m:
        return [("near_person", m.group(1))]
    m = re.fullmatch(r"远离(" + "|".join(NAMES) + r")", clause)
    if m:
        return [("far_person", m.group(1))]
    raise ValueError(f"Unparsed A1 clause: {clause!r}")


def parse_day(row):
    cells = [row[i].value for i in range(6)]
    loop = 1 if cells[0] == "一周目" else 2 if cells[0] == "二周目" else None
    if loop is None:
        return None
    day = int(re.search(r"Day\s*(\d+)", cells[1]).group(1))
    if day == 7:
        return None
    roster = [name for name in NAMES if name in (cells[2] or "")]
    rules = []
    notes = []
    for line in (cells[4] or "").splitlines():
        line = line.strip().rstrip("。;； ")
        if not line:
            continue
        match = re.match(r"^(" + "|".join(NAMES) + r")[：:]\s*(.*)", line)
        if not match:
            if line.startswith("另需"):
                notes.append(line)
                continue
            raise ValueError(f"Unparsed A1 row {row[0].row}: {line!r}")
        actor, text = match.groups()
        if actor not in roster:
            raise ValueError(f"A1 row {row[0].row} has rule for absent actor {actor}")
        text = re.sub(r"（[^）]*）", "", text)
        for part in re.split(r"[，,；;]", text):
            part = part.strip().removeprefix("且")
            for kind, target in parse_clause(part):
                rules.append({"actor": actor, "kind": kind, "target": target,
                              "source": part})
    if set(roster) != {r["actor"] for r in rules}:
        raise ValueError(f"A1 row {row[0].row} is missing an employee rule")
    return {"row": row[0].row, "loop": loop, "day": day,
            "roster": roster, "rules": rules, "notes": notes}


def distance(a, b):
    return abs(COORDS[a][0] - COORDS[b][0]) + abs(COORDS[a][1] - COORDS[b][1])


def satisfied(rule, places, zones, near_distance=2, far_distance=3):
    seat = places[rule["actor"]]
    kind, target = rule["kind"], rule["target"]
    if kind == "seat":
        return seat == SEATS.index(target)
    if kind == "zone":
        return seat in zones[target]
    other = SEATS.index(target) if target in SEATS else places[target]
    gap = distance(seat, other)
    if kind == "adjacent":
        return gap == 1
    if kind == "not_adjacent_seat":
        return gap != 1
    if kind in ("near_seat", "near_person"):
        return gap <= near_distance
    if kind == "far_person":
        return gap >= far_distance
    raise ValueError(kind)


def solve(day, zones, near_distance=2, far_distance=3):
    available = tuple(i for i in range(8) if not (day["loop"] == 1
                       and day["day"] >= 2 and i == 6))
    count = 0
    example = None
    for permutation in itertools.permutations(available, len(day["roster"])):
        places = dict(zip(day["roster"], permutation))
        if day["loop"] == 1 and day["day"] == 1 and places["SISY"] != 6:
            continue
        if all(satisfied(rule, places, zones, near_distance, far_distance)
               for rule in day["rules"]):
            count += 1
            if example is None:
                example = {actor: SEATS[seat] for actor, seat in places.items()}
    return {"count": count, "example": example}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("Docs/DDayABA1Audit.json"))
    args = parser.parse_args()
    workbook = openpyxl.load_workbook(args.workbook, data_only=False, read_only=False)
    assert workbook.sheetnames == ["A1-座位通关", "A2-即时线索",
                                  "B1-每日任务与剧情", "B2-固定事件"]
    sheet = workbook.worksheets[0]
    days = [parsed for row in sheet.iter_rows(min_row=4)
            if (parsed := parse_day(row)) is not None]
    if len(days) != 12:
        raise ValueError(f"Expected 12 seating days, found {len(days)}")
    output = {"source": str(args.workbook),
              "source_sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
              "grid": {name: COORDS[i] for i, name in enumerate(SEATS)},
              "semantics": {"adjacent": "Manhattan distance 1", "near_person": "Manhattan <=2",
                            "far_person": "Manhattan >=3", "seatB3": 6,
                            "first_loop_day2_to_day6_B3": "unavailable and empty"},
              "baseline_zones": {k: [SEATS[i] for i in v] for k, v in BASELINE.items()},
              "tuned_zones": {k: [SEATS[i] for i in v] for k, v in TUNED.items()},
              "days": []}
    for day in days:
        baseline = solve(day, BASELINE)
        tuned = solve(day, TUNED)
        output["days"].append({**day, "baseline": baseline, "tuned": tuned})
        print(f"L{day['loop']} D{day['day']}: baseline={baseline['count']}, "
              f"tuned={tuned['count']}; example={tuned['example']}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
